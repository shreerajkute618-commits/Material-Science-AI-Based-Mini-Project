"""
Automated Document Generator for Project Report (DOCX).
Converts the structured academic report into a professionally formatted Word document.
"""

import os
import sys
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

import config

def set_cell_background(cell, fill_hex):
    """Sets background shading of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tc_pr.append(shd)

def generate_report_docx():
    """Generates a professional academic project report DOCX."""
    doc = Document()

    # Document Title / Cover Page
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run("AI-BASED IDENTIFICATION AND SORTING OF RECYCLABLE MATERIALS\n")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 58, 138)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run("A Computer Vision & Deep Learning System for Materials Science & Circular Engineering\n\n")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_meta = meta_p.add_run(
        "DEPARTMENT OF MATERIALS SCIENCE AND ENGINEERING\n"
        "[COLLEGE / UNIVERSITY NAME - EDITABLE]\n"
        "ACADEMIC YEAR: 2026 – 2027\n\n"
    )
    run_meta.font.name = "Calibri"
    run_meta.font.size = Pt(12)
    run_meta.font.bold = True

    # Student metadata table
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_info = [
        ("Project Title", "AI-Based Identification and Sorting of Recyclable Materials"),
        ("Student Name", "[STUDENT NAME - EDITABLE]"),
        ("Roll / Reg. No.", "[ROLL NUMBER - EDITABLE]"),
        ("Faculty Supervisor", "[SUPERVISOR NAME - EDITABLE]"),
        ("Date of Submission", "October 2026")
    ]
    for i, (k, v) in enumerate(meta_info):
        cell_k, cell_v = table.cell(i, 0), table.cell(i, 1)
        cell_k.text, cell_v.text = k, v
        cell_k.paragraphs[0].runs[0].font.bold = True
        set_cell_background(cell_k, "F1F5F9")
        set_cell_background(cell_v, "FFFFFF")

    doc.add_page_break()

    # Section 1: Abstract
    doc.add_heading("ABSTRACT", level=1)
    p_abs = doc.add_paragraph(
        "Municipal solid waste management represents one of the most critical environmental and materials "
        "lifecycle bottlenecks of modern civilization. With global municipal waste projected to reach 3.4 billion "
        "metric tons annually by 2050, traditional manual sorting techniques in Materials Recovery Facilities (MRFs) "
        "exhibit low throughput, high operational expenses, and significant occupational health hazards.\n\n"
        "This project develops an end-to-end, automated, vision-based classification and simulated sorting system "
        "specifically grounded in Materials Science and Engineering principles. Using Transfer Learning with a "
        "lightweight MobileNetV2 convolutional neural network, material samples are identified from digital imagery "
        "and live optical camera feeds across six standard circular economy categories: Plastic, Metal, Paper, Cardboard, "
        "Glass, and Other / Non-recyclable waste. The system achieves 75.78% empirical test accuracy with zero data leakage "
        "on the Stanford TrashNet benchmark, mapping visual features to fundamental polymer, metallic, and ceramic microstructures."
    )
    p_abs.paragraph_format.line_spacing = 1.25

    # Section 2: Introduction & Problem Statement
    doc.add_heading("1. INTRODUCTION & PROBLEM STATEMENT", level=1)
    doc.add_paragraph(
        "Materials Science and Engineering (MSE) investigates the relationships between the synthesis, structure, "
        "properties, and performance of engineering materials. In modern industrial paradigms, materials engineering "
        "extends beyond initial manufacturing to encompass the complete circular lifecycle—addressing end-of-life recovery, "
        "metallurgical remelting, polymer depolymerization, and sustainable material stewardship.\n\n"
        "Manual waste picking suffers from heavy cross-contamination (>20%), severe occupational hazards (sharp metals, "
        "bioaerosols, toxic leachates), high labor costs, and a complete lack of digital compositional auditing."
    )

    # Section 3: Objectives
    doc.add_heading("2. PROJECT OBJECTIVES", level=1)
    objectives = [
        "Train and validate a MobileNetV2 deep CNN classifier for 6 recyclable material classes.",
        "Ensure rigorous zero data leakage using an empirical 70/15/15 train/val/test split.",
        "Link vision classifications with thermodynamic, chemical, and physical material properties.",
        "Implement uncertainty detection guardrails (threshold = 50%) for ambiguous samples.",
        "Develop an interactive Streamlit web dashboard with live camera capture and CSV history export.",
        "Create a simulated sorting twin demonstrating pneumatic and eddy-current chute routing."
    ]
    for obj in objectives:
        doc.add_paragraph(obj, style='List Bullet')

    # Section 4: Materials Science Integration
    doc.add_heading("3. MATERIALS SCIENCE CONCEPTS & ANALYSIS", level=1)
    mat_table = doc.add_table(rows=7, cols=4)
    mat_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Material Class", "Microstructure & Bonding", "Key Physical Properties", "Recycling Mechanism"]
    for j, h in enumerate(headers):
        cell = mat_table.cell(0, j)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        set_cell_background(cell, "E2E8F0")

    mat_rows = [
        ("Plastic", "Long-chain macromolecular polymers with covalent backbone & van der Waals forces.", "Density: 0.91-1.45 g/cm³; Melting: 110-260°C.", "Mechanical recycling (sink-float, shredding, extrusion); 4-6 cycle limit."),
        ("Metal", "Crystalline lattice (FCC/BCC) with delocalized metallic bonding sea.", "Density: 2.7-7.8 g/cm³; High conductivity & ductility.", "Infinitely recyclable; saves 95% energy vs. virgin bauxite smelting."),
        ("Paper", "Cellulose microfibrils with inter-chain hydrogen bonds.", "Hydrophilic; density: 0.6-0.9 g/cm³.", "Hydrapulping, flotation de-inking; 5-7 cycle fiber shortening limit."),
        ("Cardboard", "Fluted corrugated sandwich composite medium.", "Burst strength: 100-300 kPa; high kinetic damping.", "Repulping, centrifugal glue removal, re-fluting."),
        ("Glass", "Non-crystalline amorphous silicate network (SiO2) with modifiers.", "Mohs: 5.5-6.5; chemically inert; softening: ~750°C.", "100% infinitely recyclable; cullet lowers furnace melting fuel usage."),
        ("Other / Trash", "Heterogeneous multi-material packaging (metallized films).", "Thermodynamically incompatible blended phases.", "Refuse-Derived Fuel (RDF) gasification or sanitary landfilling.")
    ]
    for i, row in enumerate(mat_rows):
        for j, val in enumerate(row):
            c = mat_table.cell(i + 1, j)
            c.text = val
            if j == 0:
                c.paragraphs[0].runs[0].font.bold = True
            set_cell_background(c, "F8FAFC" if i % 2 == 0 else "FFFFFF")

    # Section 5: Experimental Evaluation Results
    doc.add_heading("4. EXPERIMENTAL EVALUATION RESULTS", level=1)
    doc.add_paragraph(
        "The model was evaluated strictly on the held-out test split of 384 images from the Stanford TrashNet benchmark.\n"
        "• Overall Test Accuracy: 75.78%\n"
        "• Weighted F1-Score: 0.7347\n"
        "• Macro F1-Score: 0.6485"
    )

    eval_table = doc.add_table(rows=7, cols=5)
    eval_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    eval_headers = ["Material Class", "Precision", "Recall", "F1-Score", "Test Samples"]
    for j, h in enumerate(eval_headers):
        cell = eval_table.cell(0, j)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        set_cell_background(cell, "DBEAFE")

    eval_rows = [
        ("Cardboard", "0.86", "0.84", "0.85", "61"),
        ("Paper", "0.78", "0.92", "0.85", "90"),
        ("Metal", "0.69", "0.81", "0.75", "62"),
        ("Plastic", "0.77", "0.73", "0.75", "73"),
        ("Glass", "0.69", "0.71", "0.70", "76"),
        ("Trash", "0.00", "0.00", "0.00", "22")
    ]
    for i, row in enumerate(eval_rows):
        for j, val in enumerate(row):
            c = eval_table.cell(i + 1, j)
            c.text = val
            if j == 0:
                c.paragraphs[0].runs[0].font.bold = True
            set_cell_background(c, "F8FAFC" if i % 2 == 0 else "FFFFFF")

    doc.add_heading("5. CONCLUSION", level=1)
    doc.add_paragraph(
        "This project successfully demonstrates an automated, materials-aware recyclable sorting system combining "
        "computer vision with sustainable circular engineering. The resulting web application and simulated sorting twin "
        "provide a robust platform for modern materials engineering laboratory education and industrial circular economy applications."
    )

    out_path = config.DOCS_DIR / "PROJECT_REPORT.docx"
    doc.save(out_path)
    print(f"[SUCCESS] Updated project report DOCX generated at {out_path}")
    return out_path


if __name__ == "__main__":
    generate_report_docx()
