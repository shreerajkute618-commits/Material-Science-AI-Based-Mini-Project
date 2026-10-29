"""
Automated PowerPoint Presentation Generator.
Generates a professional 10-slide PowerPoint (.pptx) presentation for project viva voce.
"""

import os
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

import config

def add_slide_header(slide, title_text, category_text="Materials Science & Circular Engineering"):
    """Adds a standard styled header and category tag to a presentation slide."""
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.4), Inches(1.1))
    tf = title_box.text_frame
    tf.word_wrap = True
    
    p_cat = tf.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.size = Pt(10)
    p_cat.font.bold = True
    p_cat.font.color.rgb = RGBColor(14, 116, 144)  # Cyan/Teal
    
    p_title = tf.add_paragraph()
    p_title.text = title_text
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(30, 41, 59)  # Slate dark

def add_bullet_points(slide, points, left=Inches(0.8), top=Inches(1.6), width=Inches(8.4), height=Inches(5.0)):
    """Adds bullet points to slide."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    
    for i, pt in enumerate(points):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = pt
        p.font.size = Pt(15)
        p.font.color.rgb = RGBColor(51, 65, 85)
        p.space_after = Pt(12)

def generate_presentation_pptx():
    """Generates the 10-slide presentation PPTX."""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Slide 1: Title Slide
    s1 = prs.slides.add_slide(blank_layout)
    tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(8.0), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "AI-BASED IDENTIFICATION AND SORTING OF RECYCLABLE MATERIALS"
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = RGBColor(30, 58, 138)
    p.space_after = Pt(14)
    
    p_sub = tf1.add_paragraph()
    p_sub.text = "Computer Vision & Deep Learning for Materials Science & Circular Engineering"
    p_sub.font.size = Pt(15)
    p_sub.font.color.rgb = RGBColor(71, 85, 105)
    p_sub.space_after = Pt(24)
    
    p_meta = tf1.add_paragraph()
    p_meta.text = (
        "Candidate: [STUDENT NAME - EDITABLE] | Roll No: [ROLL NUMBER - EDITABLE]\n"
        "Department of Materials Science and Engineering\n"
        "[COLLEGE / UNIVERSITY NAME - EDITABLE]\n"
        "Academic Year: 2026 – 2027"
    )
    p_meta.font.size = Pt(13)
    p_meta.font.color.rgb = RGBColor(100, 116, 139)

    # Slide 2: Problem Statement
    s2 = prs.slides.add_slide(blank_layout)
    add_slide_header(s2, "The Global Waste Challenge & Problem Statement")
    add_bullet_points(s2, [
        "• Global municipal solid waste generation is projected to exceed 3.4 billion tonnes by 2050.",
        "• Material Recovery Facilities (MRFs) depend heavily on hazardous manual picking lines.",
        "• High Cross-Contamination: Human sorting fatigue leads to >20% impurity in recycled scrap bales.",
        "• Occupational Hazards: Workers face sharp metals, broken glass, airborne pathogens, and chemical leachates.",
        "• The Need: An automated, high-speed, vision-driven classification and sorting system."
    ])

    # Slide 3: Project Objectives
    s3 = prs.slides.add_slide(blank_layout)
    add_slide_header(s3, "Project Objectives & Scope")
    add_bullet_points(s3, [
        "• Deep Learning Classification: Categorize waste into 6 classes (Plastic, Metal, Paper, Cardboard, Glass, Trash).",
        "• Materials Science Integration: Link visual features to macromolecular, crystalline, and amorphous structures.",
        "• Zero Data Leakage: Guarantee unbiased evaluation using an empirical 70/15/15 train/val/test split.",
        "• Uncertainty Guardrails: Flag images below 50% confidence to prevent overconfident erroneous sorting.",
        "• Digital Twin Simulation: Provide a real-time automated conveyor and chute diverter demonstration.",
        "• User Interface & Auditing: Build an interactive Streamlit app with webcam input and CSV history export."
    ])

    # Slide 4: Materials Science Integration
    s4 = prs.slides.add_slide(blank_layout)
    add_slide_header(s4, "Materials Science & Circular Engineering Foundations")
    add_bullet_points(s4, [
        "• Polymers (Plastics): Long covalent chains; mechanical recycling limited to 4-6 cycles due to thermal degradation.",
        "• Metals (Al-Alloys & Steels): Crystalline lattices with metallic bonding; infinitely recyclable without property loss.",
        "• Paper & Cardboard: Hydrophilic cellulose webs; mechanical pulping shortens fibers, requiring eventual downcycling.",
        "• Glass (Silicates): Amorphous SiO2 network; cullet addition reduces furnace energy by 2-3% per 10% added.",
        "• Other / Refuse: Non-separable multi-layer laminates diverted to Refuse-Derived Fuel (RDF) gasification."
    ])

    # Slide 5: System Architecture
    s5 = prs.slides.add_slide(blank_layout)
    add_slide_header(s5, "End-to-End System Architecture")
    add_bullet_points(s5, [
        "• Input Layer: RGB Image ingestion from user file upload or live browser webcam feed.",
        "• Preprocessing: EXIF rotation correction, resizing to 224x224 pixels, and ImageNet normalization.",
        "• CNN Backbone: Pretrained MobileNetV2 feature extractor leveraging inverted residual blocks.",
        "• Classifier Head: Global Average Pooling -> Dropout (0.2) -> Linear (128) -> ReLU -> Dropout (0.3) -> Linear (6).",
        "• Softmax & Decision Engine: Probability calibration, uncertainty evaluation, and materials knowledge lookup.",
        "• Actuation & Logging: Visual conveyor chute triggering and persistent CSV audit logging."
    ])

    # Slide 6: Deep Learning Methodology
    s6 = prs.slides.add_slide(blank_layout)
    add_slide_header(s6, "Deep Learning Architecture: Why MobileNetV2?")
    add_bullet_points(s6, [
        "• Depthwise Separable Convolutions: Splits spatial filtering from cross-channel combination, reducing computation by 8-9x.",
        "• Inverted Residuals & Linear Bottlenecks: Expands to high dimensions for feature extraction, then projects back linearly.",
        "• Parameter Efficiency: ~3.4 million parameters compared to 138 million for VGG16, preventing edge latency bottlenecks.",
        "• Transfer Learning: Exploits rich hierarchical visual features pretrained on 1.4 million ImageNet samples."
    ])

    # Slide 7: Dataset & Rigorous Split
    s7 = prs.slides.add_slide(blank_layout)
    add_slide_header(s7, "Benchmark Dataset & Stratified Splitting")
    add_bullet_points(s7, [
        "• Benchmark Dataset: Stanford TrashNet (2,527 genuine real-world waste images across 6 classes).",
        "• Training Set (70% - 1,766 images): Trained with random rotation, flips, and color jitter augmentation.",
        "• Validation Set (15% - 377 images): Used for learning rate scheduling and early model checkpointing.",
        "• Test Set (15% - 384 images): Held out strictly for final unbiased empirical evaluation.",
        "• Absolute Zero Leakage: Fixed random seeds ensure completely independent partitions."
    ])

    # Slide 8: Experimental Evaluation & Results
    s8 = prs.slides.add_slide(blank_layout)
    add_slide_header(s8, "Experimental Evaluation & Empirical Findings")
    add_bullet_points(s8, [
        "• Multi-Metric Evaluation: Overall Accuracy, Macro F1-Score, Weighted Recall, and Confusion Matrix.",
        "• High Performance Classes: Cardboard, Glass, and Metal achieve high precision due to distinct specular and edge signatures.",
        "• Ambiguity Handling: Flexible plastics occasionally resemble paper; caught effectively by the 50% uncertainty filter.",
        "• Full Reproducibility: All evaluation metrics, confusion matrix heatmaps, and loss graphs are automatically generated."
    ])

    # Slide 9: Automated Sorting Simulation & Web UI
    s9 = prs.slides.add_slide(blank_layout)
    add_slide_header(s9, "Interactive Demonstration & Digital Twin")
    add_bullet_points(s9, [
        "• Interactive Streamlit App: Clean, responsive user experience for laboratory demonstrations.",
        "• Real-Time Webcam Classification: Live camera stream integration for on-the-spot physical testing.",
        "• Confidence Distribution Bar Chart: Transparent visualization of all 6 class probability scores.",
        "• Conveyor & Chute Diverter Twin: Animated representation of industrial optical and magnetic sorters.",
        "• CSV Export: Instant download of complete prediction history for auditing and compliance."
    ])

    # Slide 10: Conclusion & Future Engineering Scope
    s10 = prs.slides.add_slide(blank_layout)
    add_slide_header(s10, "Conclusions & Future Engineering Scope")
    add_bullet_points(s10, [
        "• Conclusion: Successfully designed, trained, and demonstrated a scientifically grounded recyclable sorting AI.",
        "• Key Takeaway: Bridges AI computer vision with thermodynamic materials properties and circular lifecycle engineering.",
        "• Future Scope 1 - Hyperspectral SWIR: Detect exact polymer absorption bands (distinguish PET, HDPE, PP).",
        "• Future Scope 2 - 3D RGB-D Vision: Compute volume and center-of-mass for robotic delta pickers.",
        "• Future Scope 3 - Edge Deployment: Deploy compiled TensorRT engine on NVIDIA Jetson embedded hardware."
    ])

    out_path = config.DOCS_DIR / "PROJECT_PRESENTATION.pptx"
    prs.save(out_path)
    print(f"[SUCCESS] PowerPoint presentation generated at {out_path}")
    return out_path


if __name__ == "__main__":
    generate_presentation_pptx()
